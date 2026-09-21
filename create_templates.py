import os
from docx import Document

# Ensure templates directory exists
templates_dir = os.path.join("backend", "modules", "document_generator", "templates")
os.makedirs(templates_dir, exist_ok=True)

# 1. General Affidavit
doc_affidavit = Document()
doc_affidavit.add_heading("AFFIDAVIT", 0)
doc_affidavit.add_paragraph("I, {{ deponent_name }}, son/daughter/wife of {{ father_name }}, aged about {{ age }} years, residing at {{ address }}, do hereby solemnly affirm and state as follows:")
doc_affidavit.add_paragraph("{{ statement }}")
doc_affidavit.add_paragraph("I hereby declare that the contents of this affidavit are true to the best of my knowledge and belief, and nothing has been concealed therein.")
doc_affidavit.add_paragraph("Date: {{ date }}\n\nDeponent Signature: ___________________")
doc_affidavit.save(os.path.join(templates_dir, "general_affidavit_template.docx"))

# 2. FIR Draft
doc_fir = Document()
doc_fir.add_heading("FIRST INFORMATION REPORT (FIR)", 0)
doc_fir.add_paragraph("To,\nThe Station House Officer (SHO),\nPolice Station: {{ police_station }},\nCity/District: {{ city }}")
doc_fir.add_paragraph("Subject: Information regarding a cognizable offense.")
doc_fir.add_paragraph("Respected Sir/Madam,\nI, {{ complainant_name }}, wish to bring to your attention the following incident that occurred on {{ incident_date }}.")
doc_fir.add_paragraph("Details of the Incident:\n{{ incident_details }}")
doc_fir.add_paragraph("Suspect Details:\n{{ suspect_details }}")
doc_fir.add_paragraph("I request you to register an FIR and take necessary legal action immediately.")
doc_fir.add_paragraph("Yours faithfully,\n\n{{ complainant_name }}")
doc_fir.save(os.path.join(templates_dir, "fir_draft_template.docx"))

# 3. Legal Notice
doc_notice = Document()
doc_notice.add_heading("LEGAL NOTICE", 0)
doc_notice.add_paragraph("From:\n{{ sender_name }}\n\nTo:\n{{ recipient_name }}\n{{ recipient_address }}")
doc_notice.add_paragraph("Subject: {{ subject }}")
doc_notice.add_paragraph("Sir/Madam,\nUnder instructions from my client, I hereby serve upon you the following Legal Notice:")
doc_notice.add_paragraph("{{ notice_body }}")
doc_notice.add_paragraph("Demand / Action Required:\n{{ demand }}")
doc_notice.add_paragraph("You are hereby called upon to comply with the aforementioned demand within {{ deadline_days }} days from the receipt of this notice, failing which my client shall be constrained to initiate legal proceedings against you entirely at your risk, cost, and consequence.")
doc_notice.add_paragraph("Yours sincerely,\n\n{{ sender_name }}")
doc_notice.save(os.path.join(templates_dir, "legal_notice_template.docx"))

# 4. Power of Attorney
doc_poa = Document()
doc_poa.add_heading("GENERAL POWER OF ATTORNEY", 0)
doc_poa.add_paragraph("KNOW ALL MEN BY THESE PRESENTS that I, {{ principal_name }}, residing at {{ principal_address }}, do hereby appoint, nominate, and constitute {{ agent_name }}, residing at {{ agent_address }}, as my true and lawful attorney.")
doc_poa.add_paragraph("My attorney shall have full power and authority to act on my behalf in all matters, including financial, legal, and administrative tasks, as I myself could do if personally present.")
doc_poa.add_paragraph("This Power of Attorney shall take effect from {{ effective_date }} and shall remain in force until explicitly revoked in writing.")
doc_poa.add_paragraph("IN WITNESS WHEREOF, I have hereunto set my hand on this ____ day of ________, 20__.")
doc_poa.add_paragraph("\n\n{{ principal_name }}\n(Principal)")
doc_poa.save(os.path.join(templates_dir, "power_of_attorney_template.docx"))

# 5. Will
doc_will = Document()
doc_will.add_heading("LAST WILL AND TESTAMENT", 0)
doc_will.add_paragraph("I, {{ testator_name }}, residing at {{ testator_address }}, being of sound mind and memory, do hereby make, publish, and declare this to be my Last Will and Testament, revoking all prior Wills and codicils.")
doc_will.add_paragraph("I hereby appoint {{ executor_name }} as the Executor of this my Last Will.")
doc_will.add_paragraph("I give, devise, and bequeath all my properties and assets in the following manner:")
doc_will.add_paragraph("{{ beneficiary_details }}")
doc_will.add_paragraph("IN WITNESS WHEREOF, I have signed this Will on this date: {{ date }}.")
doc_will.add_paragraph("\n\n{{ testator_name }}\n(Testator)")
doc_will.save(os.path.join(templates_dir, "will_template.docx"))

# 6. Employment Agreement
doc_emp = Document()
doc_emp.add_heading("EMPLOYMENT AGREEMENT", 0)
doc_emp.add_paragraph("This Employment Agreement is made and entered into on {{ start_date }}, by and between {{ employer_name }} (hereinafter referred to as the 'Employer') and {{ employee_name }} (hereinafter referred to as the 'Employee').")
doc_emp.add_paragraph("1. POSITION: The Employer agrees to employ the Employee in the position of {{ job_title }}.")
doc_emp.add_paragraph("2. COMPENSATION: The Employer shall pay the Employee a monthly salary of ₹{{ salary }}.")
doc_emp.add_paragraph("3. NOTICE PERIOD: Either party may terminate this agreement by providing a written notice of {{ notice_period }} month(s).")
doc_emp.add_paragraph("IN WITNESS WHEREOF, the parties hereto have executed this Agreement.")
doc_emp.add_paragraph("\n\n___________________\t\t___________________\nEmployer Signature\t\tEmployee Signature")
doc_emp.save(os.path.join(templates_dir, "employment_agreement_template.docx"))

print("Successfully created document templates.")
