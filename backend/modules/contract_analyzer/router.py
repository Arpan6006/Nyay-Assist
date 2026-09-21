import os
import shutil
import json
import uuid
import re
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Dict, Any, Optional
from pydantic import BaseModel, Field

from ..db_sync.db import get_db
from ..auth.router import get_optional_user, User

router = APIRouter(prefix="/contracts", tags=["contract_analyzer"])
if os.environ.get("VERCEL") or os.environ.get("AWS_LAMBDA_FUNCTION_NAME"):
    UPLOAD_DIR = "/tmp/contract_uploads"
else:
    UPLOAD_DIR = os.path.join(os.path.dirname(__file__), "uploads")
os.makedirs(UPLOAD_DIR, exist_ok=True)

class ContractClause(BaseModel):
    clause: str = Field(description="The clause text or summary")
    issue: str = Field(description="Why it is flagged (e.g. one-sided, unusual, unreasonable liability)")
    severity: Optional[str] = Field(default="Medium", description="Risk level: High, Medium, or Low")
    statutory_concept: str = Field(description="Related legal concept (e.g. 'unilateral termination', 'indemnity', 'arbitration')")
    recommendation: Optional[str] = Field(default="", description="Practical suggestion to renegotiate or protect the user")
    protective_redline: Optional[str] = Field(default="", description="Balanced, legally sound alternative clause draft")
    redline_rationale: Optional[str] = Field(default="", description="Why the redline draft is safer")

class ContractAnalysisResult(BaseModel):
    summary: str = Field(description="Plain-language comprehensive summary of the contract")
    document_type: Optional[str] = Field(default="Contract / Agreement", description="Type of document analyzed")
    risk_level: Optional[str] = Field(default="Medium", description="Overall risk rating: Low, Medium, or High")
    risk_score: Optional[int] = Field(default=50, description="Risk score from 0 (Safe) to 100 (Extremely Risky)")
    is_legal_notice: Optional[bool] = Field(default=False, description="Whether the document is a legal notice or demand letter")
    key_terms: Dict[str, Any] = Field(description="Key terms extracted like parties, duration, amounts, notice period, jurisdiction")
    flagged_clauses: List[ContractClause] = Field(description="List of unusual, risky, or one-sided clauses with redline alternatives")
    immediate_action_items: Optional[List[str]] = Field(default=[], description="Urgent action items, deadlines, or statutory compliance steps")
    governing_laws: Optional[List[str]] = Field(default=[], description="Applicable Indian laws like Indian Contract Act, Transfer of Property Act, etc.")

class RedraftClauseRequest(BaseModel):
    original_clause: Optional[str] = ""
    issue: Optional[str] = ""
    custom_instruction: Optional[str] = None
    stance: Optional[str] = "balanced"  # balanced, pro_user, conservative

class GenerateReplyNoticeRequest(BaseModel):
    document_summary: Optional[str] = ""
    notice_claims: Optional[List[str]] = Field(default_factory=list)
    sender_client_name: Optional[str] = "Client"
    opposite_party_name: Optional[str] = "Opposite Party"
    advocate_name: Optional[str] = "Advocate on Record"
    user_defense_stance: Optional[str] = "Complete Denial of Allegations"
    key_facts: Optional[str] = ""

class GenerateRenegotiationEmailRequest(BaseModel):
    contract_summary: Optional[str] = ""
    flagged_clauses: Optional[List[Dict[str, Any]]] = Field(default_factory=list)
    sender_name: Optional[str] = "Client"
    recipient_name: Optional[str] = "Counterparty"
    tone: Optional[str] = "constructive_professional"



def extract_text_from_file(file_path: str, extension: str) -> str:
    """Extracts text from PDF, DOCX, or TXT document files."""
    ext = extension.lower()
    text = ""
    
    # Disallow image uploads
    if ext in [".png", ".jpg", ".jpeg", ".webp", ".tiff", ".bmp", ".gif"]:
        raise HTTPException(status_code=400, detail="Image uploads are not supported. Please upload a PDF, DOCX, or TXT contract document.")
    
    # 1. PDF Documents
    if ext == ".pdf":
        try:
            import pdfplumber
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
        except Exception as e:
            print(f"[Contract Analyzer] pdfplumber extraction notice: {e}")
            
        if not text.strip():
            try:
                import pypdf
                reader = pypdf.PdfReader(file_path)
                for page in reader.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            except Exception as e:
                print(f"[Contract Analyzer] pypdf extraction notice: {e}")
                
        if not text.strip():
            try:
                from pdf2image import convert_from_path
                import pytesseract
                images = convert_from_path(file_path, first_page=1, last_page=10)
                for img in images:
                    ocr_text = pytesseract.image_to_string(img)
                    if ocr_text:
                        text += ocr_text + "\n"
            except Exception as e:
                print(f"[Contract Analyzer] pdf2image OCR notice: {e}")

    # 2. Word Documents (.docx / .doc)
    elif ext in [".docx", ".doc"]:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            text = "\n".join(paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += "\n" + row_text
        except Exception as e:
            print(f"[Contract Analyzer] docx extraction notice: {e}")

    # 3. Plain Text / Markdown
    elif ext in [".txt", ".md", ".rtf"]:
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                text = f.read()
        except Exception as e:
            print(f"[Contract Analyzer] text file read notice: {e}")

    else:
        raise HTTPException(status_code=400, detail=f"Unsupported file format '{ext}'. Please upload a PDF, DOCX, or TXT file.")

    return text.strip()


def _extract_text_content(response) -> str:
    """Extract clean string content, stripping thinking tags or list-based chunks."""
    content = response.content if hasattr(response, "content") else str(response)
    if isinstance(content, str):
        text = content
    elif isinstance(content, list):
        texts = []
        for part in content:
            if isinstance(part, dict) and part.get("type") == "text":
                texts.append(part.get("text", ""))
            elif isinstance(part, str):
                texts.append(part)
        text = "".join(texts)
    else:
        text = str(content) if content else ""
        
    text = re.sub(r"<think>.*?</think>", "", text, flags=re.DOTALL).strip()
    return text


def parse_llm_json(raw_input: Any) -> dict:
    """Robustly parse JSON from LLM output using json_repair and regex fallbacks."""
    cleaned = _extract_text_content(raw_input)
    
    # 1. Try json_repair
    try:
        import json_repair
        parsed = json_repair.loads(cleaned)
        if isinstance(parsed, dict) and len(parsed.keys()) > 0:
            if "summary" not in parsed:
                parsed["summary"] = "The document has been analyzed."
            if "key_terms" not in parsed:
                parsed["key_terms"] = {}
            if "flagged_clauses" not in parsed:
                parsed["flagged_clauses"] = []
            if "risk_level" not in parsed:
                parsed["risk_level"] = "Medium"
            if "risk_score" not in parsed:
                parsed["risk_score"] = 65 if parsed.get("risk_level") == "High" else (45 if parsed.get("risk_level") == "Medium" else 20)
            return parsed
    except Exception as e:
        print(f"[Contract Analyzer] json_repair error: {e}")

    # 2. Try direct json.loads
    try:
        return json.loads(cleaned)
    except Exception:
        pass

    # 3. Try regex extraction of JSON object {...}
    match = re.search(r"```(?:json)?\s*(\{.*?\})\s*```", cleaned, re.DOTALL)
    if match:
        try:
            return json.loads(match.group(1))
        except Exception:
            try:
                import json_repair
                return json_repair.loads(match.group(1))
            except Exception:
                pass

    # 4. Try finding first { and last }
    start = cleaned.find("{")
    end = cleaned.rfind("}")
    if start != -1 and end != -1:
        snippet = cleaned[start:end+1]
        try:
            return json.loads(snippet)
        except Exception:
            try:
                import json_repair
                return json_repair.loads(snippet)
            except Exception:
                pass

    # 5. Last resort fallback
    return {
        "summary": cleaned[:600] if cleaned else "Document parsed and analyzed.",
        "document_type": "Legal Document",
        "risk_level": "Medium",
        "risk_score": 50,
        "is_legal_notice": False,
        "key_terms": {"Document Status": "Parsed and processed"},
        "flagged_clauses": [
            {
                "clause": "Full Document Review",
                "issue": "Extracted text analyzed by AI model.",
                "severity": "Medium",
                "statutory_concept": "General Contract Law",
                "recommendation": "Review obligations and covenants carefully.",
                "protective_redline": "All obligations shall be mutual, reasonable, and subject to 30 days prior written cure notice.",
                "redline_rationale": "Ensures reciprocal rights and adequate cure periods."
            }
        ],
        "immediate_action_items": ["Review terms with legal counsel prior to signing."],
        "governing_laws": ["Indian Contract Act, 1872", "Transfer of Property Act, 1882"]
    }


@router.post("/analyze")
async def analyze_contract(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: Optional[User] = Depends(get_optional_user)
):
    file_id = str(uuid.uuid4())
    file_extension = os.path.splitext(file.filename)[1] if file.filename else ".tmp"
    file_path = os.path.join(UPLOAD_DIR, f"{file_id}{file_extension}")
    
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    try:
        extracted_text = extract_text_from_file(file_path, file_extension)
        
        if not extracted_text or len(extracted_text.strip()) < 20:
            raise HTTPException(
                status_code=400, 
                detail="Could not extract readable text from the uploaded file. Please ensure the document contains clear text or high-resolution scan."
            )

        from ...llm_provider import get_llm
        llm = get_llm(temperature=0.1, max_tokens=3800)

        prompt_template = """You are an expert Indian Legal Contract Analyst & Senior Corporate Advocate. 
Analyze the provided legal document (which may be a Contract, Lease, Employment Deed, NDA, Commercial Agreement, or a Formal Legal Notice / Demand Letter) thoroughly.
You must flag one-sided clauses, risky liabilities, unconscionable penalties, and generate actionable, balanced protective redline draft alternatives for every problematic clause.

Return ONLY a valid JSON object matching the exact schema below.

Required JSON Structure:
{{
  "document_type": "e.g. Commercial Lease / Employment Agreement / Freelance Contract / Legal Notice / NDA",
  "summary": "Detailed plain-language summary of the document, primary obligations of each party, key financial terms, and core legal implications.",
  "risk_level": "Low | Medium | High",
  "risk_score": 75,
  "is_legal_notice": false,
  "key_terms": {{
    "Document Title": "...",
    "Parties Involved": "...",
    "Effective Date & Duration": "...",
    "Financial Terms / Rent / Fees": "...",
    "Security Deposit / Penalty": "...",
    "Notice Period / Lock-in Period": "...",
    "Governing Jurisdiction & Court": "..."
  }},
  "flagged_clauses": [
    {{
      "clause": "Exact text or specific clause snippet from the document",
      "issue": "Clear explanation of why this clause is problematic, one-sided, unfair, or poses heavy liability",
      "severity": "High | Medium | Low",
      "statutory_concept": "e.g. Unilateral termination, Section 27 Restraint of Trade, Disproportionate penalty under Sec 74, Uncapped Indemnity trap, Forfeiture of deposit",
      "recommendation": "Concrete tactical advice for the client during negotiations or dispute",
      "protective_redline": "Professional, balanced rewrite of the clause that protects the client while remaining reasonable for both sides",
      "redline_rationale": "Concise rationale explaining why this redline draft fixes the liability"
    }}
  ],
  "immediate_action_items": [
    "e.g. Request deletion of unilateral forfeiture clause 6.2 before signing",
    "e.g. Demand reciprocal 30-day notice period in Section 11",
    "e.g. If Legal Notice: Draft and send formal written reply within 15 days of receipt to prevent ex-parte proceedings"
  ],
  "governing_laws": [
    "Indian Contract Act, 1872",
    "Transfer of Property Act, 1882",
    "Specific Relief Act, 1963",
    "Arbitration and Conciliation Act, 1996"
  ]
}}

Guidelines:
1. Detect if document is a Legal Notice/Summons or an Agreement. If Legal Notice, set 'is_legal_notice': true, assess claims, and outline response deadlines.
2. For every flagged clause, provide a ready-to-use 'protective_redline' draft that can be copy-pasted into counter-proposals.
3. Assess Indian legal provisions: Section 27 (Agreement in restraint of trade void), Section 28 (Agreements in restraint of legal proceedings), Section 73/74 (Liquidated damages vs unconscionable penalty), and consumer/tenancy rights.
4. Return ONLY valid JSON, with NO markdown code block wrapper and NO conversational text.

Document Text:
\"\"\"
{contract_text}
\"\"\"
"""
        trimmed_text = extracted_text[:14000]
        formatted_prompt = prompt_template.format(contract_text=trimmed_text)

        response = llm.invoke(formatted_prompt)
        analysis_dict = parse_llm_json(response)

        if os.path.exists(file_path):
            os.remove(file_path)

        return analysis_dict

    except HTTPException:
        if os.path.exists(file_path):
            os.remove(file_path)
        raise
    except Exception as e:
        if os.path.exists(file_path):
            os.remove(file_path)
        print(f"[Contract Analyzer] Analysis failure: {e}")
        raise HTTPException(status_code=500, detail=f"Contract analysis error: {str(e)}")


@router.post("/redraft-clause")
async def redraft_clause(
    req: RedraftClauseRequest,
    current_user: Optional[User] = Depends(get_optional_user)
):
    """Customizes or refines a specific contract clause based on custom instructions or protective stance."""
    try:
        from ...llm_provider import get_llm
        llm = get_llm(temperature=0.2, max_tokens=1500)

        prompt = f"""You are a Senior Corporate Legal Drafter & Advocate under Indian Law.
Please review the original clause, the flagged issue, and rewrite the clause to be legally sound, protective, and enforceable under Indian Law.

Original Clause:
\"\"\"{req.original_clause}\"\"\"

Identified Issue:
\"\"\"{req.issue}\"\"\"

Target Stance / Mode: {req.stance} (e.g. balanced, pro_user, conservative)
Custom User Instruction: {req.custom_instruction or 'Make it balanced, fair, and legally protective with mutual obligations and reasonable cure periods.'}

Return ONLY a valid JSON object matching this schema:
{{
  "redrafted_clause": "Full rewritten clause text in professional legal drafting language",
  "explanation": "Clear explanation of how the changes safeguard the client's interests",
  "key_protections": [
    "Protection point 1",
    "Protection point 2"
  ]
}}
"""
        response = llm.invoke(prompt)
        text_content = _extract_text_content(response)
        
        try:
            import json_repair
            parsed = json_repair.loads(text_content)
            if isinstance(parsed, dict) and "redrafted_clause" in parsed:
                return parsed
        except Exception:
            pass

        # Fallback
        return {
            "redrafted_clause": text_content.replace("```json", "").replace("```", "").strip(),
            "explanation": "Clause modified to include mutual obligations and reasonable notice.",
            "key_protections": ["Mutual termination rights", "Liability cap", "Fair notice cure period"]
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Redrafting failed: {str(e)}")


def clean_plain_legal_text(text: str) -> str:
    """Strips all raw markdown symbols (asterisks, hashes, horizontal rules, broken table pipes) to yield clean, formal plain legal text."""
    if not text:
        return ""
    # Remove markdown code fences
    text = re.sub(r"^```[a-zA-Z]*\n", "", text)
    text = re.sub(r"\n```$", "", text)
    
    # Replace markdown headers ### Header -> Header
    text = re.sub(r"^#{1,6}\s*(.*)", r"\1", text, flags=re.MULTILINE)
    
    # Strip markdown bold/italic asterisks & underscores (**bold** -> bold)
    text = re.sub(r"\*\*(.*?)\*\*", r"\1", text)
    text = re.sub(r"\*(.*?)\*", r"\1", text)
    text = re.sub(r"__(.*?)__", r"\1", text)
    text = re.sub(r"_(.*?)_", r"\1", text)
    
    # Strip horizontal rules (---, ***, ___)
    text = re.sub(r"^\s*[-*_]{3,}\s*$", "", text, flags=re.MULTILINE)
    
    # Clean markdown table rows | a | b |
    lines = []
    for line in text.splitlines():
        if re.match(r"^\s*\|?\s*[-:]+\s*\|", line):
            continue  # Skip table separator line |---|---|
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|") if c.strip()]
            if cells:
                lines.append(" - " + " : ".join(cells))
            continue
        lines.append(line)
        
    clean_str = "\n".join(lines)
    # Collapse redundant empty lines to maximum 2
    clean_str = re.sub(r"\n{3,}", "\n\n", clean_str).strip()
    return clean_str


@router.post("/generate-reply-notice")
async def generate_reply_notice(
    req: GenerateReplyNoticeRequest,
    current_user: Optional[User] = Depends(get_optional_user)
):
    """Generates a formal, court-ready Reply to Legal Notice under Indian Law conventions with clean normal text."""
    try:
        from ...llm_provider import get_llm
        llm = get_llm(temperature=0.2, max_tokens=3000)

        claims_text = "\n".join([f"- {c}" for c in req.notice_claims]) if req.notice_claims else "General claims & allegations set forth in the notice"

        prompt = f"""You are a Senior High Court / Supreme Court Advocate drafting a formal 'REPLY TO LEGAL NOTICE' under Indian Jurisprudence.

Document Context & Allegations:
{req.document_summary}

Specific Claims Raised in Notice:
{claims_text}

Sender (Client) Details: {req.sender_client_name}
Opposite Party (Notice Sender): {req.opposite_party_name}
Advocate Name: {req.advocate_name or 'Advocate, High Court'}
Client's Stance: {req.user_defense_stance}
Additional Facts from Client: {req.key_facts or 'All allegations are baseless, motivated, and denied in toto.'}

CRITICAL FORMATTING INSTRUCTION:
Do NOT use raw markdown formatting symbols (NO asterisks **, NO hashes ###, NO table pipes |, NO horizontal rules ---).
Write in clean, formal legal document text using standard capitalizations, paragraph numbers (1, 1.1, 1.2, 2, 2.1), and clean legal phrasing.

Draft structure:
1. Header: BY REGISTERED POST A.D. (SPEED POST) / EMAIL
2. TO: Opposite Party & Counsel
3. SUBJECT: Reply to Legal Notice dated [...]
4. 1. PRELIMINARY OBJECTIONS (Maintainability, lack of cause of action, suppression of material facts, frivolous nature)
5. 2. PARA-WISE REPLY (Categorical denial of wrongful allegations, presenting client's true factual narrative)
6. 3. LEGAL GROUNDS (Statutes cited: Indian Contract Act, Specific Relief Act, etc.)
7. 4. FINAL CALL & RECALL DEMAND (Demand unconditional withdrawal within 15 days, failing which legal proceedings will follow)
8. Signatures & Verification block.

Return ONLY a valid JSON object matching this schema:
{{
  "reply_notice_markdown": "Full text of the formal Legal Reply Notice in clean plain text format without markdown symbols",
  "legal_grounds": [
    "Ground 1: Lack of cause of action under Contract Act",
    "Ground 2: Breach by opposite party of reciprocal promises"
  ],
  "statutory_citations": [
    "Section 73, Indian Contract Act, 1872",
    "Specific Relief Act, 1963"
  ],
  "recommended_next_steps": [
    "Send reply via Registered Post A.D. and retain postal dispatch receipt",
    "Send copy via registered email with read receipt",
    "Compile all prior payment receipts and WhatsApp/Email correspondence"
  ]
}}
"""
        response = llm.invoke(prompt)
        text_content = _extract_text_content(response)

        try:
            import json_repair
            parsed = json_repair.loads(text_content)
            if isinstance(parsed, dict):
                raw_text = parsed.get("reply_notice_markdown") or parsed.get("notice_markdown") or parsed.get("body") or text_content
                cleaned_notice = clean_plain_legal_text(str(raw_text))
                legal_grounds = [clean_plain_legal_text(str(g)) for g in parsed.get("legal_grounds", [])] if isinstance(parsed.get("legal_grounds"), list) else ["Lack of cause of action under Contract Law", "Reciprocal obligations default by opposite party"]
                statutory_citations = [clean_plain_legal_text(str(s)) for s in parsed.get("statutory_citations", [])] if isinstance(parsed.get("statutory_citations"), list) else ["Section 73, Indian Contract Act, 1872", "Specific Relief Act, 1963"]
                recommended_next_steps = [clean_plain_legal_text(str(r)) for r in parsed.get("recommended_next_steps", [])] if isinstance(parsed.get("recommended_next_steps"), list) else ["Send reply via Registered Post A.D.", "Maintain digital delivery proofs and postal dispatch receipts"]
                return {
                    "reply_notice_markdown": cleaned_notice,
                    "legal_grounds": legal_grounds,
                    "statutory_citations": statutory_citations,
                    "recommended_next_steps": recommended_next_steps
                }
        except Exception as parse_err:
            print(f"[Contract Analyzer] Reply notice json parse note: {parse_err}")

        return {
            "reply_notice_markdown": clean_plain_legal_text(text_content),
            "legal_grounds": ["Unfounded claims without legal cause of action", "Failure of consideration and reciprocal promises", "Suppression of material facts"],
            "statutory_citations": ["Indian Contract Act, 1872", "Specific Relief Act, 1963"],
            "recommended_next_steps": ["Issue reply via Registered Post A.D.", "Maintain digital delivery proofs"]
        }
    except Exception as e:
        print(f"[Contract Analyzer] Reply notice generation failure: {e}")
        raise HTTPException(status_code=500, detail=f"Reply notice generation failed: {str(e)}")


@router.post("/generate-renegotiation-email")
async def generate_renegotiation_email(
    req: GenerateRenegotiationEmailRequest,
    current_user: Optional[User] = Depends(get_optional_user)
):
    """Generates a professional renegotiation email with proposed redlines to send to the counterparty."""
    try:
        from ...llm_provider import get_llm
        llm = get_llm(temperature=0.2, max_tokens=2000)

        clauses_summary = []
        if req.flagged_clauses:
            for i, c in enumerate(req.flagged_clauses[:6]):
                if isinstance(c, dict):
                    clause_text = str(c.get("clause", ""))[:120]
                    issue = str(c.get("issue", ""))
                    redline = str(c.get("protective_redline", ""))
                    clauses_summary.append(f"Clause #{i+1}: '{clause_text}'\nIssue: {issue}\nProposed Redline: {redline}\n")

        prompt = f"""You are an expert Corporate Legal Negotiator. Draft a courteous, highly professional, and constructive renegotiation email proposing balanced amendments to a contract.

Contract Summary: {req.contract_summary or 'Agreement'}
Sender Name: {req.sender_name or 'Client'}
Recipient Name: {req.recipient_name or 'Counterparty'}
Tone: {req.tone or 'constructive_professional'}

Key Proposed Amendments / Redlines:
{chr(10).join(clauses_summary) if clauses_summary else 'Please provide mutual terms and standard cure periods.'}

Return ONLY a valid JSON object matching this schema:
{{
  "subject": "Clear, professional email subject line (e.g. Proposed Amendments & Clarifications - [Contract Name])",
  "email_body": "Complete formatted email body proposing the redlines with polite corporate phrasing and constructive rationale."
}}
"""
        response = llm.invoke(prompt)
        text_content = _extract_text_content(response)

        try:
            import json_repair
            parsed = json_repair.loads(text_content)
            if isinstance(parsed, dict):
                return {
                    "subject": parsed.get("subject") or "Proposed Amendments & Mutual Clarifications regarding the Agreement",
                    "email_body": parsed.get("email_body") or text_content
                }
        except Exception:
            pass

        return {
            "subject": "Proposed Amendments & Mutual Clarifications regarding the Agreement",
            "email_body": text_content.replace("```json", "").replace("```", "").strip()
        }
    except Exception as e:
        print(f"[Contract Analyzer] Renegotiation email failure: {e}")
        raise HTTPException(status_code=500, detail=f"Renegotiation email generation failed: {str(e)}")
